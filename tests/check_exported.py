#!/usr/bin/env python3
"""Open the file produced by the dashboard's Excel export with openpyxl and
confirm it is a genuine, intact workbook — the sheets, merged ranges and fill
colours of untouched cells must all survive, and the edits must be present.

Run tests/verify.mjs first; it writes .verify-out/exported.xlsx.

    python3 tests/check_exported.py
"""
from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
ORIGINAL = REPO / "source" / "Weekly_Project_Plan.xlsx"
EXPORTED = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO / ".verify-out" / "exported.xlsx"

failures = 0


def check(name: str, ok: bool, detail: str = "") -> None:
    global failures
    print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f" — {detail}" if detail and not ok else ""))
    if not ok:
        failures += 1


def fills(ws):
    out = {}
    for row in ws.iter_rows():
        for cell in row:
            f = cell.fill
            if f is not None and f.fill_type == "solid":
                rgb = getattr(f.start_color, "rgb", None)
                if isinstance(rgb, str):
                    out[cell.coordinate] = rgb
    return out


def main() -> int:
    if not EXPORTED.exists():
        print(f"Exported workbook not found at {EXPORTED}.\nRun:  node tests/verify.mjs")
        return 2

    print(f"Checking {EXPORTED.relative_to(REPO)}")

    check("file is a valid ZIP", zipfile.is_zipfile(EXPORTED))
    with zipfile.ZipFile(EXPORTED) as zf:
        check("ZIP reports no corrupt members", zf.testzip() is None)
        check("[Content_Types].xml present", "[Content_Types].xml" in zf.namelist())

    original = openpyxl.load_workbook(ORIGINAL)
    exported = openpyxl.load_workbook(EXPORTED)

    check("openpyxl opens the exported workbook", True)
    check(
        "same sheets in the same order",
        exported.sheetnames == original.sheetnames,
        f"{exported.sheetnames} vs {original.sheetnames}",
    )

    for name in original.sheetnames:
        o, e = original[name], exported[name]
        o_merges = sorted(str(r) for r in o.merged_cells.ranges)
        e_merges = sorted(str(r) for r in e.merged_cells.ranges)
        check(f"{name}: merged ranges preserved", o_merges == e_merges,
              f"{len(o_merges)} vs {len(e_merges)}")

        o_fills, e_fills = fills(o), fills(e)
        check(f"{name}: fill colours preserved", o_fills == e_fills,
              f"{len(o_fills)} vs {len(e_fills)} coloured cells")

        o_widths = {k: v.width for k, v in o.column_dimensions.items()}
        e_widths = {k: v.width for k, v in e.column_dimensions.items()}
        check(f"{name}: column widths preserved", o_widths == e_widths)

    # The edits verify.mjs actually made, re-read through an independent parser.
    edits_file = EXPORTED.parent / "edits.json"
    if not edits_file.exists():
        print(f"  edits.json not found next to {EXPORTED.name}. Run:  node tests/verify.mjs")
        return 2
    expected = {
        (e["sheet"], e["ref"]): e["value"]
        for e in json.loads(edits_file.read_text(encoding="utf-8"))
    }
    check("edit list handed over by verify.mjs", bool(expected), "empty edit list")

    for (sheet, ref), want in expected.items():
        got = exported[sheet][ref].value
        check(f"edit landed in {sheet}!{ref}", got == want, f"expected {want!r}, got {got!r}")

    # Everything else must be untouched.
    drift = []
    for name in original.sheetnames:
        o, e = original[name], exported[name]
        refs = {c.coordinate for row in o.iter_rows() for c in row if c.value is not None}
        refs |= {c.coordinate for row in e.iter_rows() for c in row if c.value is not None}
        for ref in refs:
            if (name, ref) in expected:
                continue
            if o[ref].value != e[ref].value:
                drift.append(f"{name}!{ref}: {o[ref].value!r} -> {e[ref].value!r}")
    check("no other cell changed", not drift, "; ".join(drift[:4]))

    print("\n" + ("ALL CHECKS PASSED" if failures == 0 else f"{failures} CHECK(S) FAILED"))
    return 0 if failures == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
